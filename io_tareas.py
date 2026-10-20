"""Plantilla, exportación e importación de tareas (Excel / CSV).

El formato de exportación es el mismo que el de la plantilla, así que un
archivo exportado se puede volver a importar.
"""
from __future__ import annotations

import io
import re
import unicodedata
from datetime import date, datetime, timedelta

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

import db

# Columna interna -> encabezado de la plantilla
ENCABEZADOS = {
    "id": "ID",
    "tarea": "Tarea",
    "estado": "Estado",
    "prioridad": "Prioridad",
    "responsable_ejecucion": "Responsable de ejecución",
    "responsable_validacion": "Responsable de validación (mail)",
    "link_formulario": "Link formulario",
    "fecha_limite": "Fecha límite (DD/MM/AAAA)",
    "observaciones": "Observaciones",
}
CAMPOS = [c for c in ENCABEZADOS if c != "id"]

INSTRUCCIONES = [
    ("ID", "No", "Dejalo vacío para tareas nuevas. Solo lo usan los archivos exportados de este tablero.", ""),
    ("Tarea", "Sí", "Texto libre.", "Conciliar pagos de octubre"),
    ("Estado", "No (Pendiente)", "Pendiente, En Progreso, En Revisión o Completado. "
     "También acepta: por hacer, en curso, a validar, cumplido, hecho, done…", "En Progreso"),
    ("Prioridad", "No (Media)", "Baja, Media, Alta o Urgente.", "Alta"),
    ("Responsable de ejecución", "No", "Nombre de quien ejecuta la tarea.", "Ana Pérez"),
    ("Responsable de validación (mail)", "No", "Mail válido. Se le puede enviar la notificación al importar.",
     "jefe@empresa.com"),
    ("Link formulario", "No", "Link al formulario de Google donde se ejecuta la tarea.",
     "https://forms.gle/abc123"),
    ("Fecha límite (DD/MM/AAAA)", "No", "Fecha de Excel o texto DD/MM/AAAA.", "31/10/2026"),
    ("Observaciones", "No", "Nota inicial para tareas nuevas (queda en el historial de observaciones). "
     "Al exportar, muestra el historial completo; al actualizar por ID se ignora.",
     "Falta el remito de la sucursal 3"),
]

_ALIAS = {
    "id": ["id"],
    "tarea": ["tarea", "nombre", "nombredelatarea", "titulo", "task", "tareas", "actividad"],
    "estado": ["estado", "status", "situacion", "lista", "columna"],
    "prioridad": ["prioridad", "priority", "urgencia"],
    "responsable_ejecucion": [
        "responsabledeejecucion", "responsableejecucion", "responsable", "asignado",
        "asignadoa", "ejecuta", "ejecutor", "owner", "assignee", "encargado",
    ],
    "responsable_validacion": [
        "responsabledevalidacion", "responsablevalidacion", "validador", "valida",
        "validacion", "mailvalidador", "emailvalidador", "mailvalidacion", "emailvalidacion",
    ],
    "link_formulario": [
        "linkformulario", "linkdelformulario", "formulario", "link", "url", "formulariogoogle",
    ],
    "fecha_limite": [
        "fechalimite", "fechadevencimiento", "vencimiento", "vence", "fecha", "duedate",
        "deadline", "fechaentrega",
    ],
    "observaciones": [
        "observaciones", "observacion", "notas", "nota", "comentarios", "comentario",
        "notes", "comments", "obs",
    ],
}
ALIAS = {a: campo for campo, lista in _ALIAS.items() for a in lista}

_ESTADOS = {
    "Pendiente": ["pendiente", "por hacer", "to do", "todo", "nuevo", "nueva", "abierto", "abierta", "sin empezar"],
    "En Progreso": ["en progreso", "progreso", "en curso", "en proceso", "haciendo", "doing",
                    "in progress", "iniciado", "iniciada"],
    "En Revisión": ["en revision", "revision", "review", "a validar", "por validar", "en validacion",
                    "para revisar", "a revisar"],
    "Completado": ["completado", "completada", "cumplido", "cumplida", "terminado", "terminada",
                   "hecho", "hecha", "done", "finalizado", "finalizada", "cerrado", "cerrada", "ok"],
}
ESTADO_ALIAS = {a: e for e, lista in _ESTADOS.items() for a in lista}
_PRIOS = {
    "Baja": ["baja", "low", "bajo"],
    "Media": ["media", "medio", "normal", "medium"],
    "Alta": ["alta", "alto", "high"],
    "Urgente": ["urgente", "urgent", "critica", "critico", "muy alta"],
}
PRIO_ALIAS = {a: p for p, lista in _PRIOS.items() for a in lista}


def _sin_acentos(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))


def _clave(header) -> str:
    """'Fecha límite (DD/MM/AAAA)' -> 'fechalimite'."""
    s = re.sub(r"\(.*?\)", "", str(header))
    return re.sub(r"[^a-z0-9]", "", _sin_acentos(s).lower())


def _texto(v) -> str:
    if v is None:
        return ""
    try:
        if pd.isna(v):
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _fecha(v):
    """Devuelve date o None; lanza ValueError si no se entiende."""
    if v is None or _texto(v) == "":
        return None
    if isinstance(v, pd.Timestamp):
        return v.date()
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)):  # número de serie de Excel
        if 20000 < float(v) < 80000:
            return date(1899, 12, 30) + timedelta(days=int(v))
        raise ValueError(f"fecha inválida: {v}")
    s = _texto(v).split(" ")[0]
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%d/%m/%y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"fecha inválida: «{_texto(v)}» (usá DD/MM/AAAA)")


def _estado(v) -> str:
    s = _sin_acentos(_texto(v)).lower()
    s = re.sub(r"\s+", " ", s)
    if not s:
        return "Pendiente"
    if s in ESTADO_ALIAS:
        return ESTADO_ALIAS[s]
    raise ValueError(f"estado desconocido: «{_texto(v)}»")


def _prioridad(v) -> str:
    s = _sin_acentos(_texto(v)).lower().strip()
    if not s:
        return "Media"
    if s in PRIO_ALIAS:
        return PRIO_ALIAS[s]
    raise ValueError(f"prioridad desconocida: «{_texto(v)}»")


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------
_AZUL = PatternFill("solid", fgColor="1F4E78")


def _hoja_tareas(wb: Workbook, filas: list[list]) -> None:
    ws = wb.active
    ws.title = "Tareas"
    headers = list(ENCABEZADOS.values())
    ws.append(headers)
    for f in filas:
        ws.append(f)
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = _AZUL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    anchos = [8, 40, 15, 12, 26, 32, 38, 16, 60]
    for i, w in enumerate(anchos):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ultima = max(len(filas) + 1, 500)
    dv_est = DataValidation(type="list", formula1='"' + ",".join(db.ESTADOS) + '"', allow_blank=True)
    dv_pri = DataValidation(type="list", formula1='"' + ",".join(db.PRIORIDADES) + '"', allow_blank=True)
    for dv in (dv_est, dv_pri):
        dv.showErrorMessage = False  # se aceptan sinónimos, se validan al importar
        ws.add_data_validation(dv)
    dv_est.add(f"C2:C{ultima}")
    dv_pri.add(f"D2:D{ultima}")
    for r in range(2, ultima + 1):
        ws.cell(row=r, column=8).number_format = "DD/MM/YYYY"
        ws.cell(row=r, column=9).alignment = Alignment(wrap_text=True, vertical="top")


def _hoja_instrucciones(wb: Workbook) -> None:
    ws = wb.create_sheet("Instrucciones")
    ws.append(["Columna", "Obligatoria", "Formato / valores", "Ejemplo"])
    for fila in INSTRUCCIONES:
        ws.append(list(fila))
    ws.append([])
    ws.append(["Cargá las tareas en la hoja «Tareas», una por fila, sin cambiar los encabezados."])
    ws.append(["Después, en el tablero: pestaña «Importar / Exportar» → subí el archivo → revisá la vista previa → Importar."])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = _AZUL
    for col, w in zip("ABCD", (32, 16, 70, 30)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")


def plantilla_xlsx() -> bytes:
    wb = Workbook()
    _hoja_tareas(wb, [])
    _hoja_instrucciones(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def historial_texto(obs: pd.DataFrame | None, task_id: int) -> str:
    """Historial de observaciones de una tarea en texto (de la más vieja a la más nueva)."""
    if obs is None or obs.empty:
        return ""
    sub = obs[obs["task_id"] == task_id].sort_values(["creado_en", "id"])
    lineas = []
    for _, o in sub.iterrows():
        cuando = pd.Timestamp(o["creado_en"]).strftime("%d/%m/%Y %H:%M") if pd.notna(o["creado_en"]) else ""
        autor = f" – {o['autor']}" if o["autor"] else ""
        lineas.append(f"[{cuando}{autor}] {o['texto']}")
    return "\n".join(lineas)


def exportar_xlsx(df: pd.DataFrame, obs: pd.DataFrame | None = None) -> bytes:
    filas = []
    for _, r in df.iterrows():
        f = r.get("fecha_limite")
        filas.append([
            int(r["id"]), r["tarea"], r["estado"], r["prioridad"],
            r.get("responsable_ejecucion", ""), r.get("responsable_validacion", ""),
            r.get("link_formulario", ""),
            f if (f is not None and not pd.isna(f)) else None,
            historial_texto(obs, int(r["id"])),
        ])
    wb = Workbook()
    _hoja_tareas(wb, filas)
    _hoja_instrucciones(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def exportar_csv(df: pd.DataFrame, obs: pd.DataFrame | None = None) -> bytes:
    out = df[db.COLUMNAS].copy()
    out["observaciones"] = out["id"].map(lambda i: historial_texto(obs, int(i)))
    out["fecha_limite"] = out["fecha_limite"].map(
        lambda f: f.strftime("%d/%m/%Y") if f is not None and not pd.isna(f) else ""
    )
    out = out.rename(columns=ENCABEZADOS)
    return out.to_csv(index=False, sep=";").encode("utf-8-sig")


# --------------------------------------------------------------------------
# Importación
# --------------------------------------------------------------------------
def leer_archivo(nombre: str, contenido: bytes) -> pd.DataFrame:
    nombre = nombre.lower()
    if nombre.endswith((".xlsx", ".xlsm", ".xls")):
        hojas = pd.ExcelFile(io.BytesIO(contenido))
        hoja = "Tareas" if "Tareas" in hojas.sheet_names else hojas.sheet_names[0]
        return hojas.parse(hoja, dtype=object)
    for enc in ("utf-8-sig", "latin-1"):
        try:
            texto = contenido.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    return pd.read_csv(io.StringIO(texto), sep=None, engine="python", dtype=str)


def mapear_columnas(columnas) -> tuple[dict, list]:
    """Devuelve ({columna_archivo: campo}, [columnas ignoradas])."""
    mapa, usados, ignoradas = {}, set(), []
    for col in columnas:
        campo = ALIAS.get(_clave(col))
        if campo and campo not in usados:
            mapa[col] = campo
            usados.add(campo)
        else:
            ignoradas.append(str(col))
    return mapa, ignoradas


def preparar(df_archivo: pd.DataFrame) -> dict:
    """Valida el archivo y devuelve filas válidas, errores y avisos."""
    mapa, ignoradas = mapear_columnas(df_archivo.columns)
    if "tarea" not in mapa.values():
        raise ValueError(
            "No encontré la columna «Tarea». Usá la plantilla o renombrá la columna "
            "con el nombre de las tareas a «Tarea»."
        )
    datos = df_archivo.rename(columns=mapa)[list(mapa.values())]
    validas, errores = [], []
    for i, row in datos.iterrows():
        fila_excel = i + 2  # encabezado en la fila 1
        crudo = {c: row.get(c) for c in mapa.values()}
        if all(_texto(v) == "" for v in crudo.values()):
            continue  # fila vacía
        try:
            valores = {
                "tarea": _texto(crudo.get("tarea")),
                "estado": _estado(crudo.get("estado")),
                "prioridad": _prioridad(crudo.get("prioridad")),
                "responsable_ejecucion": _texto(crudo.get("responsable_ejecucion")),
                "responsable_validacion": _texto(crudo.get("responsable_validacion")),
                "link_formulario": _texto(crudo.get("link_formulario")),
                "fecha_limite": _fecha(crudo.get("fecha_limite")),
            }
            valores = db._validate(valores)
            valores["observaciones"] = _texto(crudo.get("observaciones"))
            rid = _texto(crudo.get("id"))
            valores["id"] = int(float(rid)) if rid else None
        except ValueError as ex:
            errores.append({"Fila": fila_excel, "Tarea": _texto(crudo.get("tarea")), "Error": str(ex)})
            continue
        valores["_fila"] = fila_excel
        validas.append(valores)
    return {"validas": validas, "errores": errores, "mapa": mapa, "ignoradas": ignoradas}
